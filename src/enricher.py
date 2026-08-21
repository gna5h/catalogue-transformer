"""Orchestrate dimension enrichment of a Step 1 eProcess output workbook.

Public API:
  get_blank_rows(wb)       → list[RowSpec]
  enrich_rows_iter(rows, force_retry)  → generator of (RowSpec, DimensionResult, bool, bool)
  apply_results(wb, rows, results)     → modified Workbook
"""

from __future__ import annotations

import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Generator

from io import BytesIO

import openpyxl
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import PatternFill

sys.path.insert(0, str(Path(__file__).parent))
sys.path.insert(0, str(Path(__file__).parent / 'adapters'))

from adapters.base import DimensionResult, fetch_url, ImageResult, prepare_image
from adapters.registry import fetch_for_brand, get_adapter
import cache as dim_cache

# ---------------------------------------------------------------------------
# Sheet / column constants (from Step 1 writer)
# ---------------------------------------------------------------------------

M10_SHEETS = ['Microwaves', 'DishWashers', 'Fridges & Freezers', 'Rangehoods', 'Cooktops', 'Ovens',
              'Washing Machine', 'Dryer']

COL_BRAND      = 1
COL_PRODUCT    = 2
COL_H          = 3
COL_W          = 4
COL_D          = 5
COL_LINK       = 6
COL_CONF       = 7   # Dimension Confidence
COL_SRC        = 8   # Dimension Source
COL_IMG        = 9   # Image (embedded thumbnail)
COL_IMG_STATUS = 10  # Image Status

_AMBER_FILL = PatternFill(start_color='FFECB3', end_color='FFECB3', fill_type='solid')
_NO_FILL    = PatternFill(fill_type=None)


# ---------------------------------------------------------------------------
# Data types
# ---------------------------------------------------------------------------

@dataclass
class RowSpec:
    sheet:   str
    row_idx: int
    brand:   str
    sku:     str
    link:    str


# ---------------------------------------------------------------------------
# Reading the Step 1 workbook
# ---------------------------------------------------------------------------

def _cell_str(ws, row: int, col: int) -> str:
    v = ws.cell(row, col).value
    return str(v).strip() if v is not None else ''


def _is_blank_dim_row(ws, row_idx: int) -> bool:
    """True if all three dimension cells (H, W, D) are empty."""
    return all(
        not _cell_str(ws, row_idx, c)
        for c in (COL_H, COL_W, COL_D)
    )


def _is_header_row(ws, row_idx: int) -> bool:
    return (
        _cell_str(ws, row_idx, COL_BRAND)   == 'Brand' and
        _cell_str(ws, row_idx, COL_PRODUCT) == 'Product'
    )


def get_blank_rows(wb: Workbook) -> list[RowSpec]:
    """
    Parse a Step 1 output workbook and return all rows whose H/W/D are blank.
    Rows that already have dimensions filled are skipped.
    """
    rows: list[RowSpec] = []

    for sheet_name in M10_SHEETS:
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        in_block = False

        for row_idx in range(1, (ws.max_row or 0) + 1):
            brand   = _cell_str(ws, row_idx, COL_BRAND)
            product = _cell_str(ws, row_idx, COL_PRODUCT)

            if _is_header_row(ws, row_idx):
                in_block = True
                continue

            # Blank row signals end of block
            if not brand and not product:
                in_block = False
                continue

            # Data row: Product (col 2) is populated and it's not a header
            if in_block and product and brand != 'Brand':
                if _is_blank_dim_row(ws, row_idx):
                    rows.append(RowSpec(
                        sheet   = sheet_name,
                        row_idx = row_idx,
                        brand   = brand,
                        sku     = product,
                        link    = _cell_str(ws, row_idx, COL_LINK),
                    ))

    return rows


def summarise_blank_rows(rows: list[RowSpec]) -> dict:
    """Group blank rows by sheet and brand for pre-run display."""
    by_sheet: dict[str, int] = {}
    by_brand: dict[str, int] = {}
    for r in rows:
        by_sheet[r.sheet] = by_sheet.get(r.sheet, 0) + 1
        by_brand[r.brand] = by_brand.get(r.brand, 0) + 1
    return {'total': len(rows), 'by_sheet': by_sheet, 'by_brand': by_brand}


# ---------------------------------------------------------------------------
# Running enrichment (generator for live progress updates)
# ---------------------------------------------------------------------------

def enrich_rows_iter(
    rows: list[RowSpec],
    force_retry: bool = False,
) -> Generator[tuple[RowSpec, DimensionResult, bool, bool], None, None]:
    """
    Yield (RowSpec, DimensionResult, dim_fetched, img_fetched) for each row.

    dim_fetched — True when a full adapter call was made (dimensions re-fetched from network)
    img_fetched — True when an image was fetched (via full adapter call or image-only path)

    The two concerns are tracked independently:
      • dim_fetched=False, img_fetched=True  → dimensions were already Resolved; only
        the page was re-fetched to extract the og:image (no dimension re-parsing).
      • dim_fetched=False, img_fetched=False → both are already complete; result served
        entirely from cache with no network activity.
    """
    for row in rows:
        if not row.link:
            yield row, DimensionResult(
                confidence='Not Found',
                reason='No product URL available',
            ), False, False
            continue

        need_dims = dim_cache.should_fetch_dims(row.link, force_retry, row.brand)
        need_img  = dim_cache.should_fetch_image(row.link, row.brand)

        if need_dims:
            # Full adapter call: fetches HTML, extracts dimensions + image together
            result = fetch_for_brand(row.brand, row.link)
            dim_cache.store_result(row.link, result, row.brand)
            yield row, result, True, True

        elif need_img:
            # Dimensions already Resolved — fetch page once more for image only;
            # dimension fields in the cached result are preserved unchanged.
            cached = dim_cache.get_cached(row.link) or DimensionResult(
                confidence='Not Found', reason='Cache miss (unexpected)')
            html, _err = fetch_url(row.link)
            if html:
                adapter = get_adapter(row.brand)
                img_url = adapter.get_image_url(html, row.link) if adapter else None
                if img_url:
                    img = prepare_image(img_url)
                else:
                    img = ImageResult(status='Not Found')
                cached.image_bytes  = img.image_bytes
                cached.image_status = img.status
            dim_cache.store_result(row.link, cached, row.brand)
            yield row, cached, False, True

        else:
            # Both dimensions and image already complete — pure cache hit
            yield row, dim_cache.get_cached(row.link) or DimensionResult(
                confidence='Not Found', reason='Cache miss (unexpected)'), False, False


# ---------------------------------------------------------------------------
# Writing results back to the workbook
# ---------------------------------------------------------------------------

def _ensure_confidence_headers(ws) -> None:
    """
    Add enrichment column headers to every header row in the sheet,
    if not already present.
    """
    from openpyxl.styles import Font
    bold = Font(bold=True)

    for row_idx in range(1, (ws.max_row or 0) + 1):
        if _is_header_row(ws, row_idx):
            if not ws.cell(row_idx, COL_CONF).value:
                cell_c = ws.cell(row_idx, COL_CONF, value='Dimension Confidence')
                cell_s = ws.cell(row_idx, COL_SRC,  value='Dimension Source')
                cell_c.font = bold
                cell_s.font = bold
            if not ws.cell(row_idx, COL_IMG_STATUS).value:
                cell_i  = ws.cell(row_idx, COL_IMG,        value='Image')
                cell_is = ws.cell(row_idx, COL_IMG_STATUS, value='Image Status')
                cell_i.font  = bold
                cell_is.font = bold


def apply_results(
    wb: Workbook,
    rows: list[RowSpec],
    results: dict[tuple[str, int], DimensionResult],
) -> Workbook:
    """
    Write enrichment results into the workbook.

    results is keyed by (sheet_name, row_idx).
    Returns the modified workbook (same object, mutated in place).
    """
    # Add header columns to every M10 sheet
    for sheet_name in M10_SHEETS:
        if sheet_name in wb.sheetnames:
            _ensure_confidence_headers(wb[sheet_name])

    # Write per-row results
    for row in rows:
        result = results.get((row.sheet, row.row_idx))
        if result is None:
            continue

        ws = wb[row.sheet]

        # Dimensions
        ws.cell(row.row_idx, COL_H).value = result.height
        ws.cell(row.row_idx, COL_W).value = result.width
        ws.cell(row.row_idx, COL_D).value = result.depth

        # Remove amber fill for resolved rows; keep for unresolved
        fill = _NO_FILL if result.confidence == 'Resolved' else _AMBER_FILL
        for col in (COL_H, COL_W, COL_D):
            ws.cell(row.row_idx, col).fill = fill

        # Confidence + Source
        ws.cell(row.row_idx, COL_CONF).value = result.confidence
        ws.cell(row.row_idx, COL_SRC ).value = result.source_url

        # Image thumbnail
        if result.image_bytes:
            xl_img = XLImage(BytesIO(result.image_bytes))
            xl_img.anchor = f'I{row.row_idx}'
            ws.add_image(xl_img)
            ws.row_dimensions[row.row_idx].height = 135
            ws.column_dimensions['I'].width = 23

        # Image Status
        ws.cell(row.row_idx, COL_IMG_STATUS).value = result.image_status or 'Not Found'

    return wb


def _write_run_info_sheet(wb: Workbook, meta: dict) -> None:
    """Write a 'Run Info' sheet with run timestamp and fetch/cache counts."""
    from datetime import datetime, timezone
    from openpyxl.styles import Font

    sheet_name = 'Run Info'
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name)

    bold = Font(bold=True)
    ws.cell(1, 1, 'Key').font   = bold
    ws.cell(1, 2, 'Value').font = bold

    counts = meta.get('confidence_counts', {})
    rows = [
        ('Run Timestamp',           datetime.now(timezone.utc).strftime('%Y-%m-%dT%H:%M:%SZ')),
        ('Force Retry',             str(meta.get('force_retry', False))),
        ('Total Rows',              meta.get('total_rows', 0)),
        ('Dimension Fresh Fetches', meta.get('dim_fetched_count', 0)),
        ('Image Fresh Fetches',     meta.get('img_fetched_count', 0)),
        ('Cache Hits',              meta.get('cache_hit_count', 0)),
        ('Resolved',                counts.get('Resolved', 0)),
        ('Needs Review',            counts.get('Needs Review', 0)),
        ('Not Found',               counts.get('Not Found', 0)),
    ]
    for i, (key, val) in enumerate(rows, start=2):
        ws.cell(i, 1, key)
        ws.cell(i, 2, val)

    ws.column_dimensions['A'].width = 18
    ws.column_dimensions['B'].width = 24


def enrich_workbook(
    input_bytes: bytes,
    output_path: Path,
    force_retry: bool = False,
    progress_cb=None,
) -> tuple[list[RowSpec], dict[tuple[str, int], DimensionResult]]:
    """
    Full pipeline: read → enrich → write.

    input_bytes : raw .xlsx content (from Streamlit upload or file read)
    output_path : where to save the updated workbook
    force_retry : re-process non-resolved cached results
    progress_cb : optional callable(done: int, total: int, row: RowSpec, result)

    Returns (rows, results) for summary display.
    """
    import io
    wb = openpyxl.load_workbook(io.BytesIO(input_bytes))

    rows   = get_blank_rows(wb)
    total  = len(rows)
    results: dict[tuple[str, int], DimensionResult] = {}
    dim_fetched_count = 0
    img_fetched_count = 0
    cache_hit_count   = 0
    confidence_counts: dict[str, int] = {'Resolved': 0, 'Needs Review': 0, 'Not Found': 0}

    for i, (row, result, dim_fetched, img_fetched) in enumerate(enrich_rows_iter(rows, force_retry)):
        results[(row.sheet, row.row_idx)] = result
        if dim_fetched:
            dim_fetched_count += 1
        if img_fetched:
            img_fetched_count += 1
        if not dim_fetched and not img_fetched:
            cache_hit_count += 1
        confidence_counts[result.confidence] = confidence_counts.get(result.confidence, 0) + 1
        if progress_cb:
            progress_cb(i + 1, total, row, result)

    apply_results(wb, rows, results)
    _write_run_info_sheet(wb, {
        'force_retry':        force_retry,
        'total_rows':         total,
        'dim_fetched_count':  dim_fetched_count,
        'img_fetched_count':  img_fetched_count,
        'cache_hit_count':    cache_hit_count,
        'confidence_counts':  confidence_counts,
    })
    wb.save(output_path)

    return rows, results
