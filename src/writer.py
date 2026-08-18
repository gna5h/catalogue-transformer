"""Write eProcess / M10 import Excel workbook."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font

# Light amber highlight for blank dimension cells
_AMBER_FILL = PatternFill(start_color='FFECB3', end_color='FFECB3', fill_type='solid')
_BOLD = Font(bold=True)

# Fixed sheet order required by M10
M10_SHEETS = ['Microwaves', 'DishWashers', 'Fridges & Freezers', 'Rangehoods', 'Cooktops', 'Ovens']

_OUT_COLS = ['Brand', 'Product', 'Height', 'Width', 'Depth', 'Link']
_DIM_COL_INDICES = {3, 4, 5}   # 1-indexed column positions for H / W / D

_UNMAPPED_COLS = [
    'Brand', 'Product Group', 'Subcategory',
    'Model / SKU', 'Product Name', 'Reason Not Mapped', 'Link',
]


def _write_block(ws, start_row: int, block_name: str, rows: list[dict]) -> int:
    """
    Write one labelled block to a worksheet.

    Layout:
      row N  : block label (column A only)
      row N+1: header row (bold)
      row N+2+: data rows; Height/Width/Depth left blank with amber fill

    Returns the next available row number.
    """
    # Block label
    ws.cell(row=start_row, column=1, value=block_name)
    start_row += 1

    # Header row
    for col_idx, header in enumerate(_OUT_COLS, 1):
        cell = ws.cell(row=start_row, column=col_idx, value=header)
        cell.font = _BOLD
    start_row += 1

    # Data rows
    for row_data in rows:
        ws.cell(row=start_row, column=1, value=row_data.get('Brand', ''))
        ws.cell(row=start_row, column=2, value=row_data.get('Product', ''))
        for col_idx in _DIM_COL_INDICES:
            cell = ws.cell(row=start_row, column=col_idx, value='')
            cell.fill = _AMBER_FILL
        ws.cell(row=start_row, column=6, value=row_data.get('Link', ''))
        start_row += 1

    return start_row


def write_output(
    mapped: dict[str, dict[str, list[dict]]],
    unmapped: list[dict],
    output_path: str | Path,
) -> None:
    """
    Write the eProcess output workbook.

    Args:
        mapped:      {sheet_name: {block_name: [row_dict, ...]}}
                     Row dict keys: Brand, Product, Link
        unmapped:    [{Brand, Product Group, Subcategory, Model / SKU,
                       Product Name, Reason Not Mapped, Link}]
        output_path: destination .xlsx path
    """
    wb = Workbook()
    wb.remove(wb.active)   # discard the default blank sheet

    # --- M10 output sheets (always created in fixed order) ---
    for sheet_name in M10_SHEETS:
        ws = wb.create_sheet(sheet_name)
        blocks = mapped.get(sheet_name, {})
        current_row = 1
        first = True

        for block_name, rows in blocks.items():
            if not rows:
                continue
            if not first:
                current_row += 1   # blank separator row between blocks
            first = False
            current_row = _write_block(ws, current_row, block_name, rows)

    # --- Unmapped - Review sheet ---
    if unmapped:
        ws_u = wb.create_sheet('Unmapped - Review')
        for col_idx, header in enumerate(_UNMAPPED_COLS, 1):
            cell = ws_u.cell(row=1, column=col_idx, value=header)
            cell.font = _BOLD
        for row_idx, row_data in enumerate(unmapped, 2):
            ws_u.cell(row=row_idx, column=1, value=row_data.get('Brand', ''))
            ws_u.cell(row=row_idx, column=2, value=row_data.get('Product Group', ''))
            ws_u.cell(row=row_idx, column=3, value=row_data.get('Subcategory', ''))
            ws_u.cell(row=row_idx, column=4, value=row_data.get('Model / SKU', ''))
            ws_u.cell(row=row_idx, column=5, value=row_data.get('Product Name', ''))
            ws_u.cell(row=row_idx, column=6, value=row_data.get('Reason Not Mapped', ''))
            ws_u.cell(row=row_idx, column=7, value=row_data.get('Link', ''))

    wb.save(output_path)
