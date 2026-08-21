"""Unit tests for enricher.py using synthetic openpyxl workbooks.

Tests get_blank_rows, apply_results, and enrich_workbook in isolation — no network calls made.
"""

import io
import pytest
from unittest.mock import patch, MagicMock
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import PatternFill

from enricher import (
    get_blank_rows,
    apply_results,
    enrich_workbook,
    RowSpec,
    COL_BRAND, COL_PRODUCT, COL_H, COL_W, COL_D, COL_LINK, COL_CONF, COL_SRC,
)
from adapters.base import DimensionResult

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

_SHEET = 'Microwaves'
_AMBER_START = 'FFECB3'


def _make_workbook(data_rows: list[dict]) -> Workbook:
    """
    Build a minimal test workbook with a single 'Microwaves' sheet.

    data_rows is a list of dicts with optional keys:
      brand, product, h, w, d, link
    A header row (Brand/Product) is always inserted at row 1.
    Data rows start at row 2.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = _SHEET

    # Header row
    ws.cell(1, COL_BRAND).value   = 'Brand'
    ws.cell(1, COL_PRODUCT).value = 'Product'

    for i, row in enumerate(data_rows, start=2):
        ws.cell(i, COL_BRAND).value   = row.get('brand', 'LG')
        ws.cell(i, COL_PRODUCT).value = row.get('product', f'SKU{i}')
        ws.cell(i, COL_H).value       = row.get('h')
        ws.cell(i, COL_W).value       = row.get('w')
        ws.cell(i, COL_D).value       = row.get('d')
        ws.cell(i, COL_LINK).value    = row.get('link', f'http://example.com/sku{i}')

    return wb


def _is_amber(cell) -> bool:
    return cell.fill.fill_type == 'solid' and cell.fill.start_color.rgb.endswith(_AMBER_START)


# ---------------------------------------------------------------------------
# get_blank_rows tests
# ---------------------------------------------------------------------------

class TestGetBlankRows:

    def test_skips_rows_with_filled_hwd(self):
        wb = _make_workbook([
            {'product': 'SKU1', 'h': '820', 'w': '595', 'd': '540'},
        ])
        rows = get_blank_rows(wb)
        assert rows == []

    def test_includes_rows_with_empty_hwd(self):
        wb = _make_workbook([
            {'product': 'SKU1'},
        ])
        rows = get_blank_rows(wb)
        assert len(rows) == 1
        assert rows[0].sku == 'SKU1'
        assert rows[0].sheet == _SHEET

    def test_skips_rows_with_no_product_sku(self):
        wb = _make_workbook([
            {'brand': 'LG', 'product': ''},
        ])
        rows = get_blank_rows(wb)
        assert rows == []

    def test_mixed_rows_only_blank_dim_rows_returned(self):
        wb = _make_workbook([
            {'product': 'SKU1', 'h': '820', 'w': '595', 'd': '540'},  # filled — skipped
            {'product': 'SKU2'},                                        # blank dims — included
            {'product': 'SKU3', 'h': '900', 'w': '600', 'd': '550'},  # filled — skipped
        ])
        rows = get_blank_rows(wb)
        assert len(rows) == 1
        assert rows[0].sku == 'SKU2'

    def test_idempotency_already_filled_rows_not_included(self):
        # Simulates calling get_blank_rows after apply_results has written values:
        # filled rows must not appear again.
        wb = _make_workbook([
            {'product': 'SKU1', 'h': '820', 'w': '595', 'd': '540'},
        ])
        rows_first  = get_blank_rows(wb)
        rows_second = get_blank_rows(wb)
        assert rows_first  == []
        assert rows_second == []


# ---------------------------------------------------------------------------
# apply_results tests
# ---------------------------------------------------------------------------

class TestApplyResults:

    def _wb_with_blank_row(self, link='http://example.com/sku1') -> tuple[Workbook, RowSpec]:
        wb = _make_workbook([{'product': 'SKU1', 'link': link}])
        row = RowSpec(sheet=_SHEET, row_idx=2, brand='LG', sku='SKU1', link=link)
        return wb, row

    def test_writes_hwd_values_to_correct_cells(self):
        wb, row = self._wb_with_blank_row()
        result = DimensionResult(height='820', width='595', depth='540', confidence='Resolved')
        apply_results(wb, [row], {(_SHEET, 2): result})
        ws = wb[_SHEET]
        assert ws.cell(2, COL_H).value == '820'
        assert ws.cell(2, COL_W).value == '595'
        assert ws.cell(2, COL_D).value == '540'

    def test_removes_amber_fill_for_resolved(self):
        wb, row = self._wb_with_blank_row()
        # Pre-colour the cells amber to simulate a previously flagged row
        amber = PatternFill(start_color='FFECB3', end_color='FFECB3', fill_type='solid')
        ws = wb[_SHEET]
        for col in (COL_H, COL_W, COL_D):
            ws.cell(2, col).fill = amber

        result = DimensionResult(height='820', width='595', depth='540', confidence='Resolved')
        apply_results(wb, [row], {(_SHEET, 2): result})

        for col in (COL_H, COL_W, COL_D):
            assert ws.cell(2, col).fill.fill_type != 'solid', \
                f'Column {col} should not have solid fill for Resolved row'

    def test_keeps_amber_fill_for_needs_review(self):
        wb, row = self._wb_with_blank_row()
        result = DimensionResult(
            height='820', width=None, depth='540',
            confidence='Needs Review',
            reason='Partial dimensions',
        )
        apply_results(wb, [row], {(_SHEET, 2): result})
        ws = wb[_SHEET]
        for col in (COL_H, COL_W, COL_D):
            assert ws.cell(2, col).fill.fill_type == 'solid', \
                f'Column {col} should have amber fill for Needs Review'

    def test_keeps_amber_fill_for_not_found(self):
        wb, row = self._wb_with_blank_row()
        result = DimensionResult(confidence='Not Found', reason='No data')
        apply_results(wb, [row], {(_SHEET, 2): result})
        ws = wb[_SHEET]
        for col in (COL_H, COL_W, COL_D):
            assert ws.cell(2, col).fill.fill_type == 'solid', \
                f'Column {col} should have amber fill for Not Found'

    def test_adds_dimension_confidence_and_source_headers(self):
        wb, row = self._wb_with_blank_row()
        result = DimensionResult(height='820', width='595', depth='540', confidence='Resolved')
        apply_results(wb, [row], {(_SHEET, 2): result})
        ws = wb[_SHEET]
        assert ws.cell(1, COL_CONF).value == 'Dimension Confidence'
        assert ws.cell(1, COL_SRC).value  == 'Dimension Source'

    def test_writes_confidence_and_source_url_to_data_row(self):
        link = 'http://example.com/sku1'
        wb, row = self._wb_with_blank_row(link=link)
        result = DimensionResult(
            height='820', width='595', depth='540',
            confidence='Resolved',
            source_url=link,
        )
        apply_results(wb, [row], {(_SHEET, 2): result})
        ws = wb[_SHEET]
        assert ws.cell(2, COL_CONF).value == 'Resolved'
        assert ws.cell(2, COL_SRC).value  == link


# ---------------------------------------------------------------------------
# Run Info sheet tests
# ---------------------------------------------------------------------------

def _wb_bytes(data_rows: list[dict]) -> bytes:
    """Serialise a workbook built by _make_workbook to raw bytes."""
    wb = _make_workbook(data_rows)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class TestRunInfoSheet:

    def test_enrich_workbook_produces_run_info_sheet(self, tmp_path):
        """enrich_workbook() must create a 'Run Info' sheet in the output workbook."""
        resolved = DimensionResult(height='454', width='261', depth='328', confidence='Resolved')
        out = tmp_path / 'out.xlsx'

        with patch('enricher.fetch_for_brand', return_value=resolved), \
             patch('enricher.dim_cache.should_fetch_dims', return_value=True), \
             patch('enricher.dim_cache.should_fetch_image', return_value=False), \
             patch('enricher.dim_cache.store_result'):
            enrich_workbook(_wb_bytes([{'product': 'SKU1', 'link': 'http://example.com/sku1'}]),
                            out, force_retry=True)

        import openpyxl as xl
        wb_out = xl.load_workbook(out)
        assert 'Run Info' in wb_out.sheetnames

        ws = wb_out['Run Info']
        keys = [ws.cell(r, 1).value for r in range(2, ws.max_row + 1)]
        assert 'Run Timestamp'           in keys
        assert 'Dimension Fresh Fetches' in keys
        assert 'Image Fresh Fetches'     in keys
        assert 'Cache Hits'              in keys
        assert 'Resolved'                in keys

    def test_fetch_and_cache_counts_are_correct(self, tmp_path):
        """Dimension/Image Fresh Fetches and Cache Hits must reflect actual fetch decisions."""
        resolved  = DimensionResult(height='454', width='261', depth='328', confidence='Resolved')
        not_found = DimensionResult(confidence='Not Found', reason='No data', image_status='Embedded')
        out = tmp_path / 'out.xlsx'

        # Row 1 (sku1): needs dim fetch  → adapter called → dim_fetched=True, img_fetched=True
        # Row 2 (sku2): dims + image done → pure cache hit → dim_fetched=False, img_fetched=False
        def _should_fetch_dims(url, force_retry):
            return 'sku1' in url

        def _get_cached(url):
            return not_found   # image_status='Embedded' so should_fetch_image returns False

        with patch('enricher.fetch_for_brand', return_value=resolved), \
             patch('enricher.dim_cache.should_fetch_dims',  side_effect=_should_fetch_dims), \
             patch('enricher.dim_cache.should_fetch_image', return_value=False), \
             patch('enricher.dim_cache.get_cached',         side_effect=_get_cached), \
             patch('enricher.dim_cache.store_result'):
            enrich_workbook(
                _wb_bytes([
                    {'product': 'SKU1', 'link': 'http://example.com/sku1'},
                    {'product': 'SKU2', 'link': 'http://example.com/sku2'},
                ]),
                out, force_retry=False,
            )

        import openpyxl as xl
        ws = xl.load_workbook(out)['Run Info']
        info = {ws.cell(r, 1).value: ws.cell(r, 2).value
                for r in range(2, ws.max_row + 1)}

        assert info['Dimension Fresh Fetches'] == 1   # sku1 adapter call
        assert info['Image Fresh Fetches']     == 1   # sku1 image attempted via adapter
        assert info['Cache Hits']              == 1   # sku2 served entirely from cache
